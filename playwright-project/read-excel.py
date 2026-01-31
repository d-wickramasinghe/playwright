import openpyxl
import json

wb = openpyxl.load_workbook('ITPM_EXCEL.xlsx')
ws = wb.active

headers = [cell for cell in next(ws.iter_rows(values_only=True))]
data = []

for row in list(ws.iter_rows(values_only=True))[1:]:
    data.append(dict(zip(headers, [str(cell) if cell is not None else '' for cell in row])))

with open('test-cases.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print(f"Exported {len(data)} test cases to test-cases.json")
